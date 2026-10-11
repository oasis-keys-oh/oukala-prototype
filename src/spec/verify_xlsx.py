import re,sys,collections
from openpyxl import load_workbook
f=sys.argv[1]
wb=load_workbook(f)
print(wb.sheetnames)
ws=wb['Requirements']
print('freeze',ws.freeze_panes,'filter',ws.auto_filter.ref,'DV',[ (str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation],'CF',len(ws.conditional_formatting))
BAN=re.compile(r'\b(FAF|Codification|ASU|ASC|FASB|GASB)\b|Omar|Hamid',re.I)
hits=0;empties=collections.Counter();ids=set();dup=0
hdr=[c.value for c in ws[1]]
for i,row in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
    if row[0] is None: break
    if row[0] in ids: dup+=1
    ids.add(row[0])
    for j,v in enumerate(row):
        if isinstance(v,str) and BAN.search(v): hits+=1; print('BAN',row[0],hdr[j],BAN.search(v).group(0))
    for j in (0,1,2,5,6,7,11,17,19,20,21):
        if row[j] in (None,''): empties[hdr[j]]+=1
print('rows',len(ids),'dup ids',dup,'empties',dict(empties),'banned hits',hits)
for name in wb.sheetnames:
    w=wb[name]
    n=0
    for r in w.iter_rows(values_only=True):
        for v in r:
            if isinstance(v,str) and BAN.search(v): n+=1
    print(name,w.max_row,w.max_column,'ban',n)
wv=load_workbook(f,data_only=True)
errs=0
for name in wv.sheetnames:
    for r in wv[name].iter_rows(values_only=True):
        for v in r:
            if isinstance(v,str) and v.startswith('#'): errs+=1
print('error cells',errs)
s=wv['Summary']
for r in s.iter_rows(min_row=4,max_row=60,values_only=True):
    if r[0]: print(r[:6])
